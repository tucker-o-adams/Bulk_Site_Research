# -*- coding: utf-8 -*-
"""Cross-check three independent estimators of existing site load."""
import openpyxl, re
from collections import defaultdict

P = r'C:\Users\tucke\OneDrive\Documents\Claude\Projects\TBDI Modeling\Mireye\TBDI Windstream Top 200 COs v6.29.26 vTA.xlsx'
wb = openpyxl.load_workbook(P, data_only=True)

# ---- 1. asset list: generators per CLLI
ws = wb['COs_Asset_List']
hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
gens = defaultdict(list)
allassets = defaultdict(list)
kinds = defaultdict(int)
for r in range(2, ws.max_row + 1):
    clli = ws.cell(r, hdr['CLLI - A']).value
    model = (ws.cell(r, hdr['Asset Model']).value or '')
    kw = ws.cell(r, hdr['KW']).value
    fuel = ws.cell(r, hdr['FUEL SYSTEM']).value
    rt = ws.cell(r, hdr['Generator Run Time (Hours)']).value
    if not clli:
        continue
    kinds[model] += 1
    allassets[clli].append((model, kw))
    if 'GENERATOR' in model.upper():
        try:
            kwv = float(kw)
        except (TypeError, ValueError):
            kwv = None
        gens[clli].append((kwv, fuel, rt))

print('Asset Model types in COs_Asset_List:')
for k, v in sorted(kinds.items(), key=lambda x: -x[1]):
    print(f'   {v:5}  {k}')
print(f'\nCLLIs with >=1 generator: {len(gens)} of {len(allassets)} CLLIs in asset list')

# ---- 2. Top 200 sheet: utilities spend, sqft, racks
ws2 = wb['Top 200 COs']
h2 = {ws2.cell(1, c).value: c for c in range(1, ws2.max_column + 1)}
PRICE = {'KY': .073, 'GA': .0684, 'OH': .0995, 'TX': .0633, 'AR': .062, 'IA': .0626}

rows = []
for r in range(2, ws2.max_row + 1):
    clli = ws2.cell(r, h2['CLLI']).value
    if not clli:
        continue
    def g(name):
        c = h2.get(name)
        return ws2.cell(r, c).value if c else None
    def f(name):
        try:
            return float(g(name))
        except (TypeError, ValueError):
            return None
    rows.append(dict(
        clli=clli, state=g('State'), sqft=f('Square Ft-Interior'), racks=f('Rack Count'),
        util25=f('KOIOS - UTILITIES 2025'), util26=f('KOIOS - UTILITIES 2026'),
        gen_kw=sum(k for k, _, _ in gens.get(clli, []) if k) or None,
        gen_n=len(gens.get(clli, [])),
    ))

print(f'\nTop 200 rows: {len(rows)}')
print(f'  with Square Ft-Interior: {sum(1 for x in rows if x["sqft"])}')
print(f'  with KOIOS UTILITIES 2025: {sum(1 for x in rows if x["util25"])}')
print(f'  with generator kW: {sum(1 for x in rows if x["gen_kw"])}')
print(f'  with rack count: {sum(1 for x in rows if x["racks"])}')

print('\n' + '=' * 100)
print('THREE ESTIMATORS COMPARED  (sqft@187.5W = TA method; util$ -> avg kW; nameplate genset kW)')
print('=' * 100)
print(f"{'CLLI':10} {'ST':3} {'sqft':>7} {'TAmw':>7} {'util25$':>10} {'avgkW':>8} {'TAmw/avgkW':>11} {'genkW':>7} {'#g':>3}")
ratios = []
for x in sorted(rows, key=lambda y: -(y['sqft'] or 0))[:28]:
    ta = (x['sqft'] * 0.1875 / 1000) if x['sqft'] else None
    pr = PRICE.get(x['state'])
    avgkw = (x['util25'] / pr / 8760) if (x['util25'] and pr) else None
    ratio = (ta * 1000 / avgkw) if (ta and avgkw) else None
    if ratio:
        ratios.append(ratio)
    print(f"{x['clli']:10} {str(x['state']):3} {x['sqft'] or 0:7.0f} "
          f"{ta or 0:7.2f} {x['util25'] or 0:10.0f} {avgkw or 0:8.0f} "
          f"{ratio or 0:11.1f} {x['gen_kw'] or 0:7.0f} {x['gen_n']:3}")

import statistics
if ratios:
    print(f'\nTA-estimate / utility-implied-avg-kW ratio: median {statistics.median(ratios):.1f}x  '
          f'range {min(ratios):.1f}x - {max(ratios):.1f}x  (n={len(ratios)})')
