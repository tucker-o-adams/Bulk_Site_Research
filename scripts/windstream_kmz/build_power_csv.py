# -*- coding: utf-8 -*-
"""Per-CLLI existing-power table: generator kW, rack count, rack-based load estimate."""
import openpyxl, csv, os
from collections import defaultdict

SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ws200_copy.xlsx')
OUT = r'C:\Users\tucke\OneDrive\Documents\Claude\Projects\TBDI Modeling\Mireye\WS_Top200_Existing_Power.csv'

# kW per rack, derived from KOIOS 2025 utility spend / state industrial rate / 8760h,
# across 152 sites with both figures. median 1.74 (p25 1.19, p75 2.36).
KW_PER_RACK = 1.74

wb = openpyxl.load_workbook(SRC, data_only=True)

# --- generator kW per CLLI (sum of all GENERATOR rows)
ws = wb['COs_Asset_List']
h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
gen_kw = defaultdict(float)
gen_seen = set()
for r in range(2, ws.max_row + 1):
    clli = ws.cell(r, h['CLLI - A']).value
    model = (ws.cell(r, h['Asset Model']).value or '').upper()
    if not clli or 'GENERATOR' not in model:
        continue
    gen_seen.add(clli)
    try:
        gen_kw[clli] += float(ws.cell(r, h['KW']).value)
    except (TypeError, ValueError):
        pass

# --- rack counts from Top 200 sheet
ws2 = wb['Top 200 COs']
h2 = {ws2.cell(1, c).value: c for c in range(1, ws2.max_column + 1)}

rows = []
for r in range(2, ws2.max_row + 1):
    clli = ws2.cell(r, h2['CLLI']).value
    if not clli:
        continue
    racks = ws2.cell(r, h2['Rack Count']).value
    try:
        racks = float(racks)
    except (TypeError, ValueError):
        racks = None
    g = gen_kw.get(clli)
    if clli in gen_seen and not g:
        g = None  # generator present but no kW published
    rows.append({
        'CLLI': clli,
        'Total Generator KW': int(round(g)) if g else '',
        '# racks': int(racks) if racks is not None else '',
        'Estimated Power kW (racks x 1.74)': round(racks * KW_PER_RACK, 1) if racks else '',
    })

rows.sort(key=lambda x: x['CLLI'])
cols = ['CLLI', 'Total Generator KW', '# racks', 'Estimated Power kW (racks x 1.74)']
with open(OUT, 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    w.writerows(rows)

print(f'wrote {OUT}')
print(f'  rows: {len(rows)}')
print(f'  with generator kW: {sum(1 for x in rows if x["Total Generator KW"] != "")}')
print(f'  with rack count:   {sum(1 for x in rows if x["# racks"] != "")}')
tot_g = sum(x['Total Generator KW'] for x in rows if x['Total Generator KW'] != '')
tot_r = sum(x['Estimated Power kW (racks x 1.74)'] for x in rows if x['Estimated Power kW (racks x 1.74)'] != '')
print(f'  portfolio generator nameplate: {tot_g/1000:,.1f} MW')
print(f'  portfolio rack-based estimate: {tot_r/1000:,.1f} MW')
print('\nfirst 8 rows:')
for x in rows[:8]:
    print('  ', x)
