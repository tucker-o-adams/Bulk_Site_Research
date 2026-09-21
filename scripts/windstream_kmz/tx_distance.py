# -*- coding: utf-8 -*-
"""Nearest transmission line per site, from the public HIFLD/EIA layer.

Method: buffer-query the polyline service around each point, return geometry in WGS84,
then compute exact point-to-segment distance locally in a local ENU frame (accurate
well under 1% at these scales). Validates against known Mireye values first.
"""
import csv, json, math, os, sys, time, urllib.parse, urllib.request

UA = {'User-Agent': 'Mozilla/5.0'}
# Full national HIFLD layer: 94,619 segments incl. 27,277 below 100 kV.
# (The FEMA Region-9 copy is filtered to >=100 kV and silently misses 69 kV lines.)
LYR = ('https://services2.arcgis.com/FiaPA4ga0iQKduv3/arcgis/rest/services/'
       'US_Electric_Power_Transmission_Lines/FeatureServer/0')
SCRATCH = os.path.dirname(os.path.abspath(__file__))

def jget(u, t=50, tries=3):
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers=UA), timeout=t) as r:
                return json.loads(r.read().decode('utf-8', 'replace'))
        except Exception as e:
            if i == tries - 1:
                return {'__error__': str(e)[:120]}
            time.sleep(1.5)

def seg_dist_m(px, py, ax, ay, bx, by):
    """Point-to-segment distance in metres, local ENU around (px,py)."""
    k = math.cos(math.radians(py))
    PX, PY = 0.0, 0.0
    AX, AY = (ax - px) * 111320 * k, (ay - py) * 110540
    BX, BY = (bx - px) * 111320 * k, (by - py) * 110540
    dx, dy = BX - AX, BY - AY
    if dx == 0 and dy == 0:
        return math.hypot(AX, AY)
    t = max(0.0, min(1.0, ((PX - AX) * dx + (PY - AY) * dy) / (dx * dx + dy * dy)))
    return math.hypot(AX + t * dx, AY + t * dy)

def nearest_line(lat, lng, radius_m=15000):
    geom = json.dumps({'x': lng, 'y': lat, 'spatialReference': {'wkid': 4326}})
    u = (LYR + '/query?f=geojson&outFields=VOLTAGE,VOLT_CLASS,OWNER,STATUS,TYPE,SOURCEDATE,'
         'ID,SUB_1,SUB_2,INFERRED'
         '&returnGeometry=true&geometryType=esriGeometryPoint&inSR=4326&outSR=4326'
         f'&distance={radius_m}&units=esriSRUnit_Meter'
         '&spatialRel=esriSpatialRelIntersects&geometry=' + urllib.parse.quote(geom))
    r = jget(u)
    if not isinstance(r, dict) or '__error__' in r or 'error' in r:
        return None, (r.get('__error__') if isinstance(r, dict) else 'err')
    best = None      # nearest line of ANY voltage
    best_hv = None   # nearest line >= 100 kV
    for f in (r.get('features') or []):
        g = f.get('geometry') or {}
        p = f.get('properties') or {}
        parts = ([g['coordinates']] if g.get('type') == 'LineString'
                 else g.get('coordinates') if g.get('type') == 'MultiLineString' else [])
        dmin = None
        for part in parts:
            for i in range(len(part) - 1):
                d = seg_dist_m(lng, lat, part[i][0], part[i][1], part[i+1][0], part[i+1][1])
                if dmin is None or d < dmin:
                    dmin = d
        if dmin is None:
            continue
        if best is None or dmin < best[0]:
            best = (dmin, p)
        try:
            kv = float(p.get('VOLTAGE'))
        except (TypeError, ValueError):
            kv = None
        if kv is not None and kv >= 100 and (best_hv is None or dmin < best_hv[0]):
            best_hv = (dmin, p)
    if best is None:
        return None, f'no line within {radius_m} m'
    return (best, best_hv), None

# ---------------- validation against known Mireye values ----------------
VALID = [
    ('NRFDOHXA', 41.31474448, -81.53703932, 893.1, 345),
    ('CRNLGA01', 34.49720847, -83.54997767, 169.5, 115),
    ('NWRKOHXA', 40.05944853, -82.40487305, 605.7, 69),
    ('ASLDKYXA', 38.4758, -82.64385, 695.8, 69),
]
print('VALIDATION vs Mireye (EIA/HIFLD source):')
print(f"{'CLLI':10} {'mine_m':>9} {'mireye_m':>9} {'diff%':>7} {'mine_kV':>8} {'mireye_kV':>9} {'owner':<28}")
for clli, la, ln, m_d, m_kv in VALID:
    res, err = nearest_line(la, ln)
    if not res:
        print(f'{clli:10} ERROR {err}'); continue
    (d, p), _hv = res
    kv = p.get('VOLTAGE')
    diff = (d - m_d) / m_d * 100 if m_d else float('nan')
    print(f"{clli:10} {d:9.1f} {m_d:9.1f} {diff:+6.1f}% {str(kv):>8} {m_kv:>9} {str(p.get('OWNER'))[:28]:<28}")

if '--validate-only' in sys.argv:
    sys.exit()

# ---------------- run all 200 ----------------
import openpyxl
wb = openpyxl.load_workbook(os.path.join(SCRATCH, 'ws200_copy.xlsx'), data_only=True)
ws = wb['Top 200 COs']
h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
rows = []
for r in range(2, ws.max_row + 1):
    clli = ws.cell(r, h['CLLI']).value
    if not clli:
        continue
    try:
        la = float(ws.cell(r, h['LAT']).value); ln = float(ws.cell(r, h['LONG']).value)
    except (TypeError, ValueError):
        continue
    rows.append((clli, ws.cell(r, h['State']).value, la, ln))

print(f'\nQuerying {len(rows)} sites...')
out = []
for i, (clli, st, la, ln) in enumerate(rows, 1):
    def ident(p):
        """Human-usable circuit name from endpoint substations."""
        a, b = (p.get('SUB_1') or '').strip(), (p.get('SUB_2') or '').strip()
        clean = lambda s: '' if (not s or s.upper().startswith(('UNKNOWN', 'NOT AVAIL'))) else s
        a, b = clean(a), clean(b)
        if a and b:
            return f'{a} - {b}'
        return a or b or ''

    def basis(p):
        try:
            v = float(p.get('VOLTAGE'))
        except (TypeError, ValueError):
            v = None
        if v is not None and v > 0:
            return 'numeric'
        vc = (p.get('VOLT_CLASS') or '').upper()
        if vc and not vc.startswith('NOT AVAIL'):
            return 'class_band'
        return 'not_published'

    res, err = nearest_line(la, ln)
    if res:
        (d, p), hv = res
        rec = dict(CLLI=clli, State=st, lat=la, lng=ln,
                   nearest_tx_line_m=round(d, 1),
                   nearest_tx_line_ft=round(d * 3.28084),
                   voltage_kv=p.get('VOLTAGE'), volt_class=p.get('VOLT_CLASS'),
                   voltage_basis=basis(p),
                   line_name=ident(p), line_from_sub=p.get('SUB_1'), line_to_sub=p.get('SUB_2'),
                   line_id=p.get('ID'), line_type=p.get('TYPE'), attrs_inferred=p.get('INFERRED'),
                   owner=p.get('OWNER'), status=p.get('STATUS'))
        if hv:
            dh, ph = hv
            rec.update(nearest_100kv_plus_m=round(dh, 1), kv_100plus=ph.get('VOLTAGE'),
                       owner_100plus=ph.get('OWNER'), line_name_100plus=ident(ph),
                       line_id_100plus=ph.get('ID'))
        else:
            rec.update(nearest_100kv_plus_m='', kv_100plus='', owner_100plus='',
                       line_name_100plus='', line_id_100plus='')
        rec['source'] = LYR
        out.append(rec)
    else:
        out.append(dict(CLLI=clli, State=st, lat=la, lng=ln,
                        nearest_tx_line_m='', nearest_tx_line_ft='', voltage_kv='',
                        volt_class='', voltage_basis='', line_name='', line_from_sub='',
                        line_to_sub='', line_id='', line_type='', attrs_inferred='',
                        owner='', status=f'NOT FOUND ({err})',
                        nearest_100kv_plus_m='', kv_100plus='', owner_100plus='',
                        line_name_100plus='', line_id_100plus='', source=LYR))
    if i % 25 == 0:
        print(f'  {i}/{len(rows)}', flush=True)

OUT = r'C:\Users\tucke\OneDrive\Documents\Claude\Projects\TBDI Modeling\Mireye\WS_Top200_Transmission_Distance.csv'
cols = ['CLLI','State','lat','lng','nearest_tx_line_m','nearest_tx_line_ft',
        'line_name','line_from_sub','line_to_sub','line_id','line_type','owner','status',
        'voltage_kv','volt_class','voltage_basis','attrs_inferred',
        'nearest_100kv_plus_m','kv_100plus','line_name_100plus','line_id_100plus','owner_100plus',
        'source']
with open(OUT, 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.DictWriter(f, fieldnames=cols); w.writeheader(); w.writerows(out)
ok = [x for x in out if x['nearest_tx_line_m'] != '']
print(f'\nwrote {OUT}\n  resolved {len(ok)}/{len(out)}')
if ok:
    ds = sorted(x['nearest_tx_line_m'] for x in ok)
    print(f'  distance m: min {ds[0]:.0f} | median {ds[len(ds)//2]:.0f} | max {ds[-1]:.0f}')
    print(f'  within 1 km: {sum(1 for d in ds if d <= 1000)} | within 3 km: {sum(1 for d in ds if d <= 3000)}')
