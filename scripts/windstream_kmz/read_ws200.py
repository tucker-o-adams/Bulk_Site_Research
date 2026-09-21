# -*- coding: utf-8 -*-
import openpyxl
p = r'C:\Users\tucke\OneDrive\Documents\Claude\Projects\TBDI Modeling\Mireye\TBDI Windstream Top 200 COs v6.29.26 vTA.xlsx'
wb = openpyxl.load_workbook(p, data_only=True)

for s in ['Top 200 COs', 'Top 50', 'COs_Asset_List', 'Sheet1']:
    ws = wb[s]
    print('=' * 78)
    print(f'SHEET: {s}  ({ws.max_row} rows x {ws.max_column} cols)')
    print('=' * 78)
    # find header row = first row with >3 non-empty cells
    hdr_row = None
    for r in range(1, min(8, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if sum(1 for v in vals if v not in (None, '')) > 3:
            hdr_row = r
            break
    hdr_row = hdr_row or 1
    hdrs = [(c, ws.cell(hdr_row, c).value) for c in range(1, ws.max_column + 1)]
    print(f'header row {hdr_row}:')
    for c, h in hdrs:
        if h not in (None, ''):
            # sample first 2 non-empty values below
            samp = []
            for r in range(hdr_row + 1, min(hdr_row + 40, ws.max_row + 1)):
                v = ws.cell(r, c).value
                if v not in (None, ''):
                    samp.append(str(v)[:26])
                if len(samp) >= 2:
                    break
            print(f'   c{c:2} | {str(h)[:38]:38} | {samp}')
    print()
