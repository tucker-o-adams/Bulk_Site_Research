# -*- coding: utf-8 -*-
"""Step 1 of broker-text extraction: dump the broker's cells, as received, into a worklist.

    .venv_fema/Scripts/python.exe scripts/bulk/extract_cells.py <column_map.json> --out Outputs/<batch>/input/

Deterministic and offline. Reads the broker workbook through a per-batch column map (drafted by Claude
at intake, checked by Tucker; format in EXTRACTION.md) and writes:

    site_list.csv   one row per site: site_id, name, section, source row
    cells.csv       one row per site x mapped cell: the text with the confidence tag split off, the tag
                    normalised to the vocabulary's scale, the dimensions the column feeds, and dup_n =
                    how many sites carry the identical text in that column (copied text is market-level,
                    not site-level evidence)
    extract_cells.json  source + column-map sha256, counts

Claude then reads cells.csv and writes broker_evidence.csv and location_clues.csv (EXTRACTION.md);
extract_check.py validates them against these cells. Nothing here interprets a cell.
"""
import argparse, csv, hashlib, json, os, re, sys
from collections import defaultdict

import openpyxl
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
VOCAB = json.load(open(os.path.join(HERE, 'extraction_vocab.json'), encoding='utf-8'))
DIMENSIONS = set(VOCAB['dimensions']) | {'location'}

CELL_COLUMNS = ['site_id', 'name', 'section', 'sheet', 'column', 'dimensions', 'source_cell', 'text', 'tag_raw', 'tag_norm', 'dup_n', 'dup_sites']


def sha(path):
    return hashlib.sha256(open(path, 'rb').read()).hexdigest()


def norm_tag(raw, levels):
    if raw is None:
        return 'untagged'
    for pattern, level in levels:
        if re.search(pattern, raw, re.I):
            return level
    raise SystemExit(f'tag [{raw}] matches no tag_levels rule in the column map; add one (levels: {VOCAB["tag_levels"]})')


def split_tag(text, pattern):
    """(body, tag or None) - the broker's trailing confidence tag, when the map defines one."""
    text = ('' if text is None else str(text)).strip()
    if pattern:
        m = re.search(pattern, text)
        if m:
            return text[:m.start()].strip(), m.group(1).strip()
    return text, None


def read_table(wb, t, cmap):
    ws = wb[t['sheet']]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next((i for i, r in enumerate(rows) if r and str(r[0] or '').strip() == t['header_first_cell']), None)
    if hdr_i is None:
        raise SystemExit(f"{t['sheet']}: no row whose first cell is {t['header_first_cell']!r}")
    hdr = [str(h).strip() if h is not None else '' for h in rows[hdr_i]]
    mapped, ignored = cmap['columns'], set(cmap.get('ignore_columns', []))
    unknown = [h for h in hdr if h and h not in mapped and h not in ignored and h not in (t['id_column'], t['name_column'])]
    if unknown:
        raise SystemExit(f'column map does not cover header(s) {unknown}: map each to dimensions or list it in ignore_columns')
    missing = [h for h in mapped if h not in hdr]
    if missing:
        raise SystemExit(f'column map names header(s) not in the sheet: {missing}')
    bad = {h: d for h, ds in mapped.items() for d in ds if d not in DIMENSIONS}
    if bad:
        raise SystemExit(f'unknown dimension(s) in column map: {bad}; see extraction_vocab.json')
    id_c, name_c = hdr.index(t['id_column']), hdr.index(t['name_column'])
    sites, cells, section = [], [], ''
    for i in range(hdr_i + 1, len(rows)):
        r = rows[i]
        sid = str(r[id_c]).strip() if r[id_c] is not None else ''
        if not sid or (t.get('skip_id_regex') and re.search(t['skip_id_regex'], sid)):
            continue
        if r[name_c] in (None, ''):                  # a banner row names the section below it
            section = sid
            continue
        sites.append({'site_id': sid, 'name': str(r[name_c]).strip(), 'section': section, 'source_row': f"{t['sheet']}!{i + 1}"})
        # the site name often carries location clues ("Oak Sub (Town / Main St & 2nd St)")
        cells.append({'site_id': sid, 'name': str(r[name_c]).strip(), 'section': section, 'sheet': t['sheet'], 'column': t['name_column'],
                      'dimensions': 'location', 'source_cell': f"{t['sheet']}!{get_column_letter(name_c + 1)}{i + 1}",
                      'text': str(r[name_c]).strip(), 'tag_raw': ''})
        for c, h in enumerate(hdr):
            if h not in mapped or not mapped[h]:
                continue
            body, tag = split_tag(r[c], cmap.get('tag_pattern'))
            cells.append({'site_id': sid, 'name': str(r[name_c]).strip(), 'section': section, 'sheet': t['sheet'], 'column': h,
                          'dimensions': ';'.join(mapped[h]), 'source_cell': f"{t['sheet']}!{get_column_letter(c + 1)}{i + 1}",
                          'text': body, 'tag_raw': tag or ''})
    return sites, cells


def read_detail(wb, d, cmap, known):
    """Key/value blocks: a title row 'SITE-ID · name' then rows of (key, value)."""
    ws = wb[d['sheet']]
    cells, current = [], None
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        k = str(r[0]).strip() if r and r[0] is not None else ''
        v = r[1] if len(r) > 1 else None
        m = re.match(d['title_regex'], k)
        if m and v in (None, ''):
            current = m.group(1) if m.group(1) in known else None
            continue
        if current and k in d['keys'] and d['keys'][k]:
            body, tag = split_tag(v, cmap.get('tag_pattern'))
            cells.append({'site_id': current, 'name': known[current]['name'], 'section': known[current]['section'], 'sheet': d['sheet'],
                          'column': k, 'dimensions': ';'.join(d['keys'][k]), 'source_cell': f"{d['sheet']}!B{i}", 'text': body, 'tag_raw': tag or ''})
    return cells


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('column_map')
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    cmap = json.load(open(a.column_map, encoding='utf-8'))
    src = cmap['source'] if os.path.isabs(cmap['source']) else os.path.join(ROOT, cmap['source'])
    wb = openpyxl.load_workbook(src, data_only=True)
    sites, cells = read_table(wb, cmap['table'], cmap)
    known = {s['site_id']: s for s in sites}
    if len(known) != len(sites):
        raise SystemExit('duplicate site ids in the table')
    for d in cmap.get('detail', []):
        bad = {k: x for k, ds in d['keys'].items() for x in ds if x not in DIMENSIONS}
        if bad:
            raise SystemExit(f'unknown dimension(s) in detail keys: {bad}')
        cells += read_detail(wb, d, cmap, known)
    levels = cmap.get('tag_levels', [])
    groups = defaultdict(list)
    for c in cells:
        c['tag_norm'] = norm_tag(c['tag_raw'] or None, levels)
        if c['text']:
            groups[(c['sheet'], c['column'], c['text'])].append(c['site_id'])
    for c in cells:
        g = groups.get((c['sheet'], c['column'], c['text']), [])
        c['dup_n'] = len(g)
        c['dup_sites'] = ' '.join(g) if len(g) > 1 else ''
    os.makedirs(a.out, exist_ok=True)
    with open(os.path.join(a.out, 'site_list.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=['site_id', 'name', 'section', 'source_row']); w.writeheader(); w.writerows(sites)
    with open(os.path.join(a.out, 'cells.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=CELL_COLUMNS); w.writeheader(); w.writerows(cells)
    info = {'source': os.path.relpath(src, ROOT), 'source_sha256': sha(src), 'column_map': os.path.relpath(os.path.abspath(a.column_map), ROOT),
            'column_map_sha256': sha(a.column_map), 'sites': len(sites), 'cells': len(cells),
            'cells_nonempty': sum(1 for c in cells if c['text']), 'cells_copied': sum(1 for c in cells if c['dup_n'] > 1)}
    json.dump(info, open(os.path.join(a.out, 'extract_cells.json'), 'w', encoding='utf-8'), indent=1)
    print(f"{len(sites)} sites, {len(cells)} cells ({info['cells_nonempty']} with text, {info['cells_copied']} copied across sites) -> {a.out}")


if __name__ == '__main__':
    main()
