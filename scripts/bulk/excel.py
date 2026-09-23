# -*- coding: utf-8 -*-
"""Write a batch's workbook from its run outputs.

    .venv_fema/Scripts/python.exe scripts/bulk/excel.py Outputs/<batch>/ [--name <title>]

Reads sites.csv, provenance.csv, run.json in the batch folder and writes
<batch>.xlsx beside them with five sheets:

    Sites        one row per site; producer bands over field names; absent cells
                 shaded grey and failed cells red so a blank is never read as zero
    Gaps         every absent / failed value with the source's note
    Sources      one row per producer: source, URL, vintage, method, fields, tallies
    Provenance   the full site x field table, filterable
    Run          input file and hash, run time, counts, rejected rows
"""
import argparse, csv, json, os, sys
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from sites import REQUIRED, OPTIONAL          # noqa: E402

BAND_TITLES = {
    'transmission': 'Transmission (HIFLD)', 'substations': 'Substations (HIFLD)', 'flood': 'Flood (FEMA NFHL)',
    'wetlands': 'Wetlands (USFWS NWI)', 'metro': 'Metro (Census urban areas)', 'datacenter': 'Data centers (PeeringDB)',
    'housing': 'Housing (Census 2020 blocks)', 'schools': 'Schools (NCES)', 'worship': 'Places of worship (HIFLD)',
    'healthcare': 'Nursing homes & hospitals (CMS)', 'parcel': 'Parcel (county / state GIS)',
    'footprint': 'Site footprint: parcel, else square around pin (FEMA, NWI)',
}
BAND_COLORS = ['1F3864', '2E5A46', '7A4A00', '4A235A', '0B5345', '6E2C00', '1B4F72', '4D5656', '5B2C6F', '145A32',
               '78281F', '1A5276', '3D3D3D']      # one per band: Site + 12 producers

FILL_ABSENT = PatternFill('solid', fgColor='E7E6E6')
FILL_FAILED = PatternFill('solid', fgColor='F8CBAD')
FILL_HDR = PatternFill('solid', fgColor='D9D9D9')
THIN = Side(style='thin', color='BFBFBF')
FONT_HDR = Font(bold=True)
FONT_BAND = Font(bold=True, color='FFFFFF')
WRAP = Alignment(wrap_text=True, vertical='top')


def num(v):
    """CSV strings back to numbers where they are numbers; booleans to True/False."""
    if v is None or v == '':
        return None
    if v in ('True', 'False'):
        return v == 'True'
    try:
        f = float(v)
        return int(f) if f.is_integer() and '.' not in v else f
    except ValueError:
        return v


def autosize(ws, min_w=8, max_w=60):
    widths = defaultdict(int)
    for row in ws.iter_rows(values_only=True):
        for i, v in enumerate(row, 1):
            if v is not None:
                widths[i] = max(widths[i], min(max_w, len(str(v))))
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = max(min_w, w + 2)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('batch')
    ap.add_argument('--name', default=None)
    a = ap.parse_args()
    b = a.batch.rstrip('/\\')
    name = a.name or os.path.basename(b)
    sites = list(csv.DictReader(open(os.path.join(b, 'sites.csv'), encoding='utf-8-sig')))
    prov = list(csv.DictReader(open(os.path.join(b, 'provenance.csv'), encoding='utf-8-sig')))
    run = json.load(open(os.path.join(b, 'run.json'), encoding='utf-8'))
    if not sites:
        raise SystemExit(f"{b}: sites.csv has no sites - every input row was rejected; see sites_rejected in run.json")
    status = {(p['site_id'], p['field']): p for p in prov}
    producers = run['producers']            # name -> {source, url, vintage, fields, status_counts}

    wb = Workbook()

    # ------------------------------------------------------------------ Sites
    ws = wb.active; ws.title = 'Sites'
    site_cols = ['site_id', 'lat', 'lng'] + [c for c in OPTIONAL if any(s.get(c) for s in sites)]
    # G: every row says which basis its answers rest on. A site with no parcel boundary is not a
    # blank row - its values are real, measured over the square around the pin (footprint
    # producer) or, in a batch run without it, at the pin.
    has_fp = any(s.get('fp_basis') for s in sites)
    has_parcel = has_fp or any('parcel_status' in s for s in sites)
    if has_parcel:
        for s in sites:
            s['basis'] = (s.get('fp_basis') or 'no footprint - rerun') if has_fp else \
                'parcel boundary' if s.get('parcel_status') == 'ok' else 'point only'
        site_cols.insert(1, 'basis')
    prod_fields = [f for p in producers.values() for f in p['fields']]
    extra_cols = [c for c in sites[0].keys() if c not in site_cols and c not in prod_fields and c not in OPTIONAL]
    bands = [('Site', site_cols + extra_cols)] + [(BAND_TITLES.get(n, n), p['fields']) for n, p in producers.items()]
    legend = (f'{name} — {len(sites)} sites — run {run["run_at"][:16].replace("T", " ")} UTC — distances in metres '
              '(1 mi = 1,609 m) — grey = source confirmed nothing there (absent), red = source failed (see Gaps)')
    if has_fp:
        legend += ('   |   basis: "parcel boundary" = a parcel polygon resolved; the fp_* footprint columns (flood zones, SFHA, '
                   'floodway, wetland acres) are measured over that parcel. "<n> m square" = no parcel resolved, so the fp_* '
                   'columns are measured over a square centred on the pin - 200 m x 200 m (9.88 ac), or sized to acres_stated '
                   'when the input gives a larger acreage - ground around the site, '
                   'not a parcel, and no parcel acreage is claimed. Point columns (fema_*, nwi_*) are always at the pin.')
    elif has_parcel:
        legend += ('   |   basis: "parcel boundary" = a parcel polygon resolved, so acreage and any parcel-clipped '
                   'figure describe the site; "point only" = no parcel service is registered for that county, so every '
                   'value is measured at the pin (flood zone at the pin, nearest wetland from the pin) and no acreage is claimed')
    ws.cell(1, 1, legend).font = Font(italic=True)
    col = 1
    for bi, (title, fields) in enumerate(bands):
        c0 = col
        for f in fields:
            ws.cell(3, col, f).font = FONT_HDR
            ws.cell(3, col).fill = FILL_HDR
            ws.cell(3, col).alignment = Alignment(wrap_text=True, vertical='bottom')
            col += 1
        ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=col - 1)
        cell = ws.cell(2, c0, title); cell.font = FONT_BAND; cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill('solid', fgColor=BAND_COLORS[bi % len(BAND_COLORS)])
    all_fields = [f for _, fs in bands for f in fs]
    for r, s in enumerate(sites, 4):
        for c, f in enumerate(all_fields, 1):
            v = num(s.get(f))
            cell = ws.cell(r, c, v)
            st = status.get((s['site_id'], f))
            if st:
                if st['status'] == 'absent':
                    cell.fill = FILL_ABSENT
                elif st['status'] == 'failed':
                    cell.fill = FILL_FAILED
            if isinstance(v, float):
                cell.number_format = '#,##0.0'
            elif isinstance(v, int) and not isinstance(v, bool) and f not in ('lat', 'lng'):
                cell.number_format = '#,##0'
    ws.freeze_panes = ws.cell(4, len(site_cols) + 1)
    ws.auto_filter.ref = f'A3:{get_column_letter(len(all_fields))}{3 + len(sites)}'
    ws.row_dimensions[3].height = 45
    autosize(ws, max_w=28)
    for c in range(1, len(all_fields) + 1):
        ws.column_dimensions[get_column_letter(c)].width = min(ws.column_dimensions[get_column_letter(c)].width, 18)

    # ------------------------------------------------------------------ Gaps
    g = wb.create_sheet('Gaps')
    g.append(['site_id', 'field', 'status', 'note', 'source'])
    for cell in g[1]:
        cell.font = FONT_HDR; cell.fill = FILL_HDR
    gaps = [p for p in prov if p['status'] in ('absent', 'failed')]
    for p in sorted(gaps, key=lambda p: (p['status'] != 'failed', p['site_id'], p['field'])):
        g.append([p['site_id'], p['field'], p['status'], p['note'], p['source']])
        if p['status'] == 'failed':
            for cell in g[g.max_row]:
                cell.fill = FILL_FAILED
    g.freeze_panes = 'A2'; g.auto_filter.ref = g.dimensions
    autosize(g, max_w=90)

    # ------------------------------------------------------------------ Sources
    so = wb.create_sheet('Sources')
    so.append(['producer', 'source', 'source_url', 'vintage', 'method', 'fields', 'ok', 'absent', 'failed', 'manual', 'note'])
    for cell in so[1]:
        cell.font = FONT_HDR; cell.fill = FILL_HDR
    by_prod = defaultdict(list)
    for p in prov:
        by_prod[p['field']].append(p)
    for n, p in producers.items():
        rows = [r for f in p['fields'] for r in by_prod.get(f, [])]
        vint = Counter(r['vintage'] for r in rows if r['vintage']).most_common(1)
        method = Counter(r['method'] for r in rows if r['method']).most_common(1)
        notes = Counter(r['note'] for r in rows if r['note'] and r['status'] == 'ok').most_common(1)
        sc = p['status_counts']
        so.append([n, p['source'], p['url'], p.get('vintage') or (vint[0][0] if vint else ''), method[0][0] if method else '',
                   ', '.join(p['fields']), sc.get('ok', 0), sc.get('absent', 0), sc.get('failed', 0), sc.get('manual', 0),
                   notes[0][0] if notes else ''])
    for row in so.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    autosize(so, max_w=70)

    # ------------------------------------------------------------------ Provenance
    pv = wb.create_sheet('Provenance')
    cols = ['site_id', 'field', 'value', 'status', 'source', 'source_url', 'vintage', 'fetched_at', 'method', 'note']
    pv.append(cols)
    for cell in pv[1]:
        cell.font = FONT_HDR; cell.fill = FILL_HDR
    for p in prov:
        pv.append([num(p['value']) if c == 'value' else p.get(c) for c in cols])
    pv.freeze_panes = 'A2'; pv.auto_filter.ref = pv.dimensions
    autosize(pv, max_w=50)

    # ------------------------------------------------------------------ Run
    rn = wb.create_sheet('Run')
    rn.append(['key', 'value'])
    for cell in rn[1]:
        cell.font = FONT_HDR; cell.fill = FILL_HDR
    for k, v in [('batch', name), ('run_at (UTC)', run['run_at']), ('elapsed_s', run['elapsed_s']),
                 ('input file', run['input']['path']), ('input sha256', run['input']['sha256']),
                 ('sites accepted', run['sites_accepted']), ('sites rejected', len(run['sites_rejected'])),
                 ('producers', ', '.join(producers)), ('provenance rows', len(prov)),
                 ('absent values', sum(1 for p in prov if p['status'] == 'absent')),
                 ('failed values', sum(1 for p in prov if p['status'] == 'failed')),
                 ('cache hits / misses', f"{run['cache']['hits']} / {run['cache']['misses']}"),
                 ('status meanings', 'ok = source returned a value; absent = source confirmed nothing there (an answer, not an error); '
                                     'failed = source unreachable, value null; manual = supplied by a person')]:
        rn.append([k, v])
    for n, sid, why in run['sites_rejected']:
        rn.append([f'rejected row {n}', f'{sid}: {why}'])
    autosize(rn, max_w=100)

    out = os.path.join(b, f'{name}.xlsx')
    wb.save(out)
    print(f'wrote {out}: Sites {len(sites)} rows x {len(all_fields)} cols | Gaps {len(gaps)} | Sources {len(producers)} | Provenance {len(prov)}')


if __name__ == '__main__':
    main()
